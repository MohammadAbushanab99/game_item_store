import csv
import io
from app.core.exceptions import ImportValidationException


def parse_import_file(
    filename: str, content: bytes
) -> tuple[list[str], list[tuple[int, dict]]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _parse_xlsx(content)
    if name.endswith(".csv"):
        return _parse_csv(content)
    raise ImportValidationException(
        "Unsupported file type. Please upload a .xlsx or .csv file."
    )


def _is_blank(value) -> bool:
    return value is None or str(value).strip() == ""


def _build_rows(raw_rows: list[list]) -> tuple[list[str], list[tuple[int, dict]]]:
    if not raw_rows:
        return ([], [])
    headers = ["" if h is None else str(h).strip().lower() for h in raw_rows[0]]
    rows: list[tuple[int, dict]] = []
    for offset, raw in enumerate(raw_rows[1:], start=2):
        if all((_is_blank(cell) for cell in raw)):
            continue
        row = {
            headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))
        }
        rows.append((offset, row))
    return (headers, rows)


def _parse_csv(content: bytes) -> tuple[list[str], list[tuple[int, dict]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportValidationException(
            "The CSV file is not valid UTF-8 text."
        ) from exc
    return _build_rows([list(r) for r in csv.reader(io.StringIO(text))])


def _parse_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, dict]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportValidationException(
            "Excel support is unavailable on the server."
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportValidationException(
            "The Excel file could not be read. Is it a valid .xlsx file?"
        ) from exc
    worksheet = workbook.active
    raw_rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
    return _build_rows(raw_rows)
